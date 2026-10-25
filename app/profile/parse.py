"""Turn uploaded files into text: resume (PDF/DOCX) and the LinkedIn data-export ZIP.

LinkedIn is only ever read from the official export the user downloads themselves;
it is never scraped.
"""

import csv
import io
import re
import zipfile

from pydantic import BaseModel

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_RESUME_CHARS = 30_000


class InputError(ValueError):
    """A user-facing problem with an uploaded file or link."""


def resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if len(data) > MAX_UPLOAD_BYTES:
        raise InputError("That file is over 10 MB.")
    if name.endswith(".pdf"):
        text = _pdf_text(data)
    elif name.endswith(".docx"):
        text = _docx_text(data)
    else:
        raise InputError("That file type isn't supported. Upload a PDF or DOCX.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 50:
        raise InputError(
            "No readable text found in that file. If it is a scanned PDF, export it "
            "again as a text PDF or DOCX."
        )
    return text[:MAX_RESUME_CHARS]


def _pdf_text(data: bytes) -> str:
    import pymupdf

    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)
    except Exception as e:  # pymupdf raises several types for broken files
        raise InputError("That PDF could not be opened. Is it password-protected?") from e


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise InputError("That DOCX could not be opened.") from e
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


class LinkedInPosition(BaseModel):
    title: str
    company: str = ""
    started: str = ""
    finished: str = ""
    description: str = ""


class LinkedInExport(BaseModel):
    headline: str = ""
    summary: str = ""
    positions: list[LinkedInPosition] = []
    skills: list[str] = []
    languages: list[dict[str, str]] = []  # {"name", "proficiency"}
    education: list[str] = []

    def as_text(self) -> str:
        lines = [self.headline, self.summary]
        for p in self.positions:
            span = f"{p.started}–{p.finished or 'present'}"
            lines.append(f"{p.title} at {p.company} ({span}). {p.description}")
        if self.skills:
            lines.append("Skills: " + ", ".join(self.skills))
        if self.languages:
            langs = ", ".join(f"{x['name']} ({x.get('proficiency', '')})" for x in self.languages)
            lines.append("Languages: " + langs)
        lines += self.education
        return "\n".join(line for line in lines if line.strip())


def _csv_rows(zf: zipfile.ZipFile, wanted: str) -> list[dict[str, str]]:
    # Exports nest files differently between versions; match on the base name.
    for info in zf.infolist():
        if info.filename.rsplit("/", 1)[-1].lower() == wanted.lower():
            raw = zf.read(info).decode("utf-8-sig", errors="replace")
            lines = raw.splitlines()
            # Some export CSVs start with a "Notes:" preamble before the header row.
            while lines and lines[0].startswith("Notes"):
                lines = lines[1:]
                while lines and lines[0].strip() and "," not in lines[0]:
                    lines = lines[1:]
                while lines and not lines[0].strip():
                    lines = lines[1:]
            return [
                {(k or "").strip(): (v or "").strip() for k, v in row.items()}
                for row in csv.DictReader(io.StringIO("\n".join(lines)))
            ]
    return []


def linkedin_export(data: bytes) -> LinkedInExport:
    if len(data) > 50 * MAX_UPLOAD_BYTES:
        raise InputError("That ZIP is over 500 MB; upload the basic export, not the full archive.")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise InputError("That isn't a ZIP file. Upload the export LinkedIn emailed you.") from e
    with zf:
        profile = _csv_rows(zf, "Profile.csv")
        positions = _csv_rows(zf, "Positions.csv")
        skills = _csv_rows(zf, "Skills.csv")
        languages = _csv_rows(zf, "Languages.csv")
        education = _csv_rows(zf, "Education.csv")
    if not (profile or positions or skills):
        raise InputError(
            "No Profile, Positions or Skills CSV found in that ZIP. Is it a LinkedIn data export?"
        )
    head = profile[0] if profile else {}
    return LinkedInExport(
        headline=head.get("Headline", ""),
        summary=head.get("Summary", ""),
        positions=[
            LinkedInPosition(
                title=p.get("Title", ""),
                company=p.get("Company Name", ""),
                started=p.get("Started On", ""),
                finished=p.get("Finished On", ""),
                description=p.get("Description", ""),
            )
            for p in positions
            if p.get("Title")
        ],
        skills=[s["Name"] for s in skills if s.get("Name")],
        languages=[
            {"name": x.get("Name", ""), "proficiency": x.get("Proficiency", "")}
            for x in languages
            if x.get("Name")
        ],
        education=[
            " ".join(filter(None, [e.get("Degree Name"), e.get("School Name")])) for e in education
        ],
    )
